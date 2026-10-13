# -*- coding: utf-8 -*-
"""
FastAPI 后端服务（多实例 + 巡访引擎）
- 后台线程轮询微信/QQ消息 → 主号线程（实例1+实例2同时巡访）
- Excel读取线程 → 小号线程 → 成功后同步给主号线程
- WebSocket 实时推送消息和任务状态到前端
"""
import asyncio
import json
import threading
import time
import platform

import cv2

from fastapi import FastAPI, WebSocket, WebSocketDisconnect, UploadFile, File
from fastapi.responses import HTMLResponse

from config import (
    POLL_INTERVAL, DEDUP_CACHE_SIZE,
    ADB_INSTANCES, ADB_PATH, EXCEL_CONFIG,
    SERVER_HOST, SERVER_PORT, UID_LENGTH,
)
from monitor import WeChatMonitor
from qq_monitor import QQMonitor
from task_handler import TaskHandler, is_uid, set_uid_length
from excel_reader import ExcelReader

# 设置UID位数
set_uid_length(UID_LENGTH)

app = FastAPI(title="群聊消息监控 - 多实例巡访")

wechat_mon = WeChatMonitor(dedup_size=DEDUP_CACHE_SIZE)
qq_mon = QQMonitor(dedup_size=DEDUP_CACHE_SIZE)

ws_clients = set()
logs = []
MAX_LOGS = 200


def add_log(text):
    entry = f"[{time.strftime('%H:%M:%S')}] {text}"
    logs.append(entry)
    if len(logs) > MAX_LOGS:
        logs.pop(0)
    broadcast({"type": "log", "data": entry})


def broadcast(message):
    msg_str = json.dumps(message, ensure_ascii=False)
    loop = getattr(app.state, "loop", None)
    if loop is None:
        return
    for ws in list(ws_clients):
        try:
            asyncio.run_coroutine_threadsafe(ws.send_text(msg_str), loop)
        except Exception:
            ws_clients.discard(ws)


def on_task_status(task):
    broadcast({"type": "task_status", "data": task_handler.task_to_dict(task)})


# ── 初始化任务处理器 ──
adb_config = {"instances": ADB_INSTANCES, "adb_path": ADB_PATH}
task_handler = TaskHandler(adb_config, excel_config=EXCEL_CONFIG,
                           on_status_change=on_task_status, on_log=add_log)

# ── 巡访引擎（仅Windows） ──
visit_engines = {}
mumu_instances = {}  # MuMuIO 实例

def init_visit_engines():
    """启动时自动发现MuMu实例并初始化巡访引擎"""
    global visit_engines, mumu_instances
    if platform.system() != "Windows":
        add_log("非Windows系统，巡访引擎不初始化（仅ADB模式）")
        return
    try:
        from mumu_discover import auto_connect_all
        # 自动发现所有MuMu实例
        discovered = auto_connect_all(on_log=add_log)
        if not discovered:
            add_log("未发现MuMu实例，请启动MuMu后点'自动检测环境'")
            return
        # 为每个发现的实例连接巡访引擎
        for name, info in discovered.items():
            _connect_engine(name, info["ipc_path"])
        task_handler.set_visit_engines(visit_engines)
    except Exception as e:
        add_log(f"巡访引擎模块加载失败: {e}")


def _connect_engine(name, ipc_config):
    """连接一个MuMu实例并创建巡访引擎"""
    global visit_engines, mumu_instances
    try:
        import sys
        if "code/fishing_v023" not in sys.path:
            sys.path.insert(0, "code/fishing_v023")
        from mumu_io import MuMuIO
        from visit_engine import VisitEngine

        # 先断开旧连接
        if name in mumu_instances:
            try:
                mumu_instances[name].close()
            except Exception:
                pass
            del mumu_instances[name]
        if name in visit_engines:
            del visit_engines[name]

        mumu = MuMuIO(ipc_config)
        mumu_instances[name] = mumu
        engine = VisitEngine(mumu, assets_dir="visit_assets", on_log=add_log)
        visit_engines[name] = engine
        task_handler.set_visit_engines(visit_engines)
        add_log(f"{name} 巡访引擎已初始化 ({mumu.width}x{mumu.height})")
        return True
    except Exception as e:
        add_log(f"{name} 巡访引擎初始化失败: {e}")
        return False


excel_reader = ExcelReader(EXCEL_CONFIG, task_handler, on_log=add_log)

monitor_running = False
monitor_thread = None


def monitor_loop():
    global monitor_running
    wechat_status = None
    qq_status = None

    while monitor_running:
        try:
            w_msgs, w_status = wechat_mon.poll()
            if w_status != wechat_status:
                wechat_status = w_status
                broadcast({"type": "connection", "platform": "微信", "status": w_status})

            for msg in w_msgs:
                hit = is_uid(msg.get("content", ""))
                broadcast({"type": "message", "data": msg, "is_uid": hit})
                if hit:
                    task_handler.submit_from_message(msg)

            q_msgs, q_status = qq_mon.poll()
            if q_status != qq_status:
                qq_status = q_status
                broadcast({"type": "connection", "platform": "QQ", "status": q_status})

            for msg in q_msgs:
                hit = is_uid(msg.get("content", ""))
                broadcast({"type": "message", "data": msg, "is_uid": hit})
                if hit:
                    task_handler.submit_from_message(msg)

        except Exception as e:
            add_log(f"监控错误: {e}")

        time.sleep(POLL_INTERVAL)


@app.on_event("startup")
async def startup():
    app.state.loop = asyncio.get_event_loop()
    task_handler.start()
    init_visit_engines()
    add_log("服务已启动（主号通道 + 小号通道）")


@app.on_event("shutdown")
async def shutdown():
    global monitor_running
    monitor_running = False
    excel_reader.stop()
    task_handler.stop()
    for mumu in mumu_instances.values():
        try:
            mumu.close()
        except Exception:
            pass
    add_log("服务已停止")


@app.get("/")
async def index():
    try:
        with open("index.html", "r", encoding="utf-8") as f:
            return HTMLResponse(f.read())
    except FileNotFoundError:
        return HTMLResponse("<h1>index.html not found</h1>")


@app.get("/api/status")
async def get_status():
    return {
        "monitor_running": monitor_running,
        "excel_running": excel_reader._running,
        "wechat": {"connected": wechat_mon._last_status == "online"},
        "qq": {"connected": qq_mon._last_status == "online"},
        "instances": task_handler.adb_manager.get_status(),
        "visit_engines": list(visit_engines.keys()),
        "stats": task_handler.get_stats(),
    }


@app.post("/api/monitor/start")
async def start_monitor():
    global monitor_running, monitor_thread
    if monitor_running:
        return {"ok": True, "msg": "已在运行"}
    monitor_running = True
    monitor_thread = threading.Thread(target=monitor_loop, daemon=True)
    monitor_thread.start()
    add_log("消息监控已启动")
    return {"ok": True}


@app.post("/api/monitor/stop")
async def stop_monitor():
    global monitor_running
    monitor_running = False
    add_log("消息监控已停止")
    return {"ok": True}


@app.post("/api/excel/start")
async def start_excel():
    if excel_reader._running:
        return {"ok": True, "msg": "已在运行"}
    excel_reader.start()
    return {"ok": True}


@app.post("/api/excel/stop")
async def stop_excel():
    excel_reader.stop()
    return {"ok": True}


@app.post("/api/adb/connect")
async def adb_connect():
    results = task_handler.adb_manager.connect_all()
    for name, ok in results.items():
        add_log(f"ADB {name} {'连接成功' if ok else '连接失败'}")
    return {"ok": all(results.values()), "results": results}


@app.post("/api/detect")
async def detect_all():
    """一键自动检测所有环境：微信/QQ窗口 + MuMu模拟器 + 巡访引擎"""
    report = {
        "wechat": {"found": False, "detail": ""},
        "qq": {"found": False, "detail": ""},
        "instances": {},
        "visit_engines": {},
        "ready": False,
    }

    # 1. 检测微信窗口
    try:
        w = wechat_mon.find_wechat_window()
        if w:
            report["wechat"]["found"] = True
            chat = wechat_mon.get_current_chat_name(w)
            report["wechat"]["detail"] = f"已检测到微信窗口" + (f"，当前聊天: {chat}" if chat else "")
            add_log(f"[检测] 微信: 已检测到窗口" + (f"，当前聊天: {chat}" if chat else ""))
        else:
            report["wechat"]["detail"] = "未检测到微信窗口，请确保微信已启动并登录"
            add_log("[检测] 微信: 未检测到")
    except Exception as e:
        report["wechat"]["detail"] = f"检测失败: {e}"
        add_log(f"[检测] 微信: 检测失败 {e}")

    # 2. 检测QQ窗口
    try:
        q = qq_mon.find_qq_window()
        if q:
            report["qq"]["found"] = True
            chat = qq_mon.get_current_chat_name(q)
            report["qq"]["detail"] = f"已检测到QQ窗口" + (f"，当前聊天: {chat}" if chat else "")
            add_log(f"[检测] QQ: 已检测到窗口" + (f"，当前聊天: {chat}" if chat else ""))
        else:
            report["qq"]["detail"] = "未检测到QQ窗口，请确保QQ已启动并登录"
            add_log("[检测] QQ: 未检测到")
    except Exception as e:
        report["qq"]["detail"] = f"检测失败: {e}"
        add_log(f"[检测] QQ: 检测失败 {e}")

    # 3. 检测每个MuMu实例（ADB + IPC）
    all_instances_ok = True
    for name, inst in task_handler.adb_manager.instances.items():
        info = {"adb": False, "adb_detail": "", "ipc": False, "ipc_detail": "", "resolution": None}

        # ADB检测
        try:
            ok = inst.ensure_connected()
            info["adb"] = ok
            if ok:
                info["adb_detail"] = f"ADB已连接 ({inst.device_serial})"
                add_log(f"[检测] {name} ADB: 已连接 {inst.device_serial}")
            else:
                info["adb_detail"] = f"ADB连接失败 ({inst.device_serial})"
                add_log(f"[检测] {name} ADB: 连接失败 {inst.device_serial}")
                all_instances_ok = False
        except Exception as e:
            info["adb_detail"] = f"ADB检测异常: {e}"
            add_log(f"[检测] {name} ADB: 异常 {e}")
            all_instances_ok = False

        # IPC/巡访引擎检测
        engine = visit_engines.get(name)
        if engine is not None:
            info["ipc"] = True
            info["resolution"] = f"{engine.width}x{engine.height}"
            loaded = len(engine._templates)
            total = len(engine._specs)
            missing = [n for n in engine._specs if n not in engine._templates]
            info["ipc_detail"] = f"巡访引擎已连接 ({engine.width}x{engine.height})，模板 {loaded}/{total}"
            if missing:
                info["ipc_detail"] += f"，缺图片: {','.join(missing)}"
            add_log(f"[检测] {name} 巡访引擎: {engine.width}x{engine.height}，模板 {loaded}/{total}")
        else:
            cfg = ADB_INSTANCES.get(name, {})
            ipc_path = cfg.get("ipc_config")
            if not ipc_path:
                info["ipc_detail"] = "未配置ipc_config，巡访引擎不可用"
            else:
                info["ipc_detail"] = f"ipc_config已配置但引擎未初始化: {ipc_path}"
            add_log(f"[检测] {name} 巡访引擎: 未初始化")
            all_instances_ok = False

        report["instances"][name] = info

    # 4. 巡访引擎汇总
    report["visit_engines"] = {
        "initialized": list(visit_engines.keys()),
        "count": len(visit_engines),
    }

    # 5. 总体就绪判断
    w_ok = report["wechat"]["found"]
    q_ok = report["qq"]["found"]
    report["ready"] = all_instances_ok and (w_ok or q_ok)
    if report["ready"]:
        add_log("[检测] 全部就绪！可以开始巡访")
    else:
        problems = []
        if not w_ok and not q_ok:
            problems.append("微信和QQ都未检测到")
        if not all_instances_ok:
            problems.append("部分模拟器实例未就绪")
        add_log(f"[检测] 尚未就绪: {'；'.join(problems)}")

    return report


@app.post("/api/adb/roles")
async def set_roles(body: dict):
    role_map = body.get("roles", {})
    task_handler.adb_manager.set_roles(role_map)
    add_log(f"角色已更新: {role_map}")
    return {"ok": True, "instances": task_handler.adb_manager.get_status()}


@app.post("/api/discover")
async def discover_mumu():
    """自动发现MuMu实例并连接巡访引擎"""
    if platform.system() != "Windows":
        return {"ok": False, "error": "仅Windows支持"}
    try:
        from mumu_discover import auto_connect_all
        discovered = auto_connect_all(on_log=add_log)
        if not discovered:
            return {"ok": False, "error": "未发现MuMu实例，请确保MuMu 12已安装并启动"}
        # 连接每个发现的实例
        connected = {}
        for name, info in discovered.items():
            ok = _connect_engine(name, info["ipc_path"])
            connected[name] = {"ok": ok, "port": info["port"], "serial": info["serial"]}
        return {"ok": True, "discovered": connected}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/api/instance/ipc")
async def set_ipc_config(body: dict):
    """动态设置实例的IPC配置路径并连接"""
    name = body.get("name")
    ipc_config = body.get("ipc_config")
    if not name or not ipc_config:
        return {"ok": False, "error": "缺少 name 或 ipc_config"}

    if platform.system() != "Windows":
        return {"ok": False, "error": "仅Windows支持"}

    ok = _connect_engine(name, ipc_config)
    return {"ok": ok}


@app.get("/api/instance/screenshot/{name}")
async def instance_screenshot(name: str):
    """获取某个实例的截图（base64），供前端显示画面"""
    if name not in visit_engines:
        return {"ok": False, "error": "巡访引擎未初始化"}
    try:
        import base64
        engine = visit_engines[name]
        frame = engine.screenshot()
        # 编码为JPEG base64
        _, buf = cv2.imencode('.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, 60])
        b64 = base64.b64encode(buf).decode('ascii')
        return {"ok": True, "image": b64, "width": engine.width, "height": engine.height}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/api/speed")
async def set_speed(body: dict):
    """设置操作间隔速度（0.16快速/0.32标准/0.65慢速）"""
    delay = float(body.get("delay", 0.32))
    for engine in visit_engines.values():
        engine.delay = delay
    add_log(f"操作速度已设置: {delay}秒")
    return {"ok": True, "delay": delay}


@app.get("/api/excel/status")
async def excel_status():
    """获取Excel UID列表和查找状态"""
    return excel_reader.get_status()


@app.post("/api/excel/upload")
async def excel_upload(file: UploadFile = File(...)):
    """上传Excel文件，自动提取UID列"""
    import os
    from openpyxl import load_workbook
    import tempfile

    # 保存上传的文件
    suffix = os.path.splitext(file.filename)[1]
    tmp_path = excel_reader.file_path
    if suffix not in ('.xlsx', '.xlsm', '.csv', '.tsv', '.txt'):
        return {"ok": False, "error": "仅支持 xlsx/xlsm/csv/tsv/txt 格式"}

    with open(tmp_path, "wb") as f:
        content = await file.read()
        f.write(content)

    # 解析UID列
    try:
        import re
        if suffix in ('.xlsx', '.xlsm'):
            wb = load_workbook(tmp_path)
            ws = wb.active
            headers = {}
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=1, column=col).value
                if val:
                    headers[str(val).strip().lower()] = col

            # 优先找uid列
            uid_col = headers.get("uid") or headers.get("id")
            has_header = uid_col is not None

            if uid_col is None:
                # 没有表头，扫描前几行找哪一列有8位数字
                uid_col = 1
                for col in range(1, ws.max_column + 1):
                    digit_count = 0
                    for row in range(1, min(ws.max_row + 1, 6)):
                        val = ws.cell(row=row, column=col).value
                        if val and re.match(r'^\d{8}$', str(val).strip()):
                            digit_count += 1
                    if digit_count > 0:
                        uid_col = col
                        break
                has_header = False

            start_row = 2 if has_header else 1
            uids = []
            for row in range(start_row, ws.max_row + 1):
                val = ws.cell(row=row, column=uid_col).value
                if val is not None:
                    uid_str = str(val).strip()
                    if uid_str:
                        uids.append(uid_str)
            wb.close()
        else:
            # CSV/TSV/TXT
            with open(tmp_path, "r", encoding="utf-8-sig") as f:
                lines = f.readlines()

            sep = '\t' if suffix == '.tsv' else ','
            rows = [l.strip().split(sep) for l in lines if l.strip()]

            if not rows:
                uids = []
            elif len(rows) == 1:
                uids = [rows[0][0].strip()] if rows[0] else []
            else:
                # 检查第一行是否有表头（第一行第一列不是8位数字）
                first_val = rows[0][0].strip() if rows[0] else ""
                has_header = not re.match(r'^\d{8}$', first_val)

                # 找UID所在列：扫描前5行，哪一列8位数字最多
                uid_col = 0
                max_digits = 0
                for col in range(len(rows[0])):
                    digit_count = 0
                    for row in rows[:5]:
                        if col < len(row) and re.match(r'^\d{8}$', row[col].strip()):
                            digit_count += 1
                    if digit_count > max_digits:
                        max_digits = digit_count
                        uid_col = col

                start = 1 if has_header else 0
                uids = []
                for row in rows[start:]:
                    if uid_col < len(row) and row[uid_col].strip():
                        uids.append(row[uid_col].strip())

        add_log(f"上传Excel: {file.filename}，提取到 {len(uids)} 个UID")
        return {"ok": True, "count": len(uids), "uids": uids[:20], "file_path": tmp_path}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.get("/api/tasks")
async def get_tasks():
    return [task_handler.task_to_dict(t) for t in task_handler.get_all_tasks()]


@app.get("/api/logs")
async def get_logs():
    return logs


@app.websocket("/ws")
async def websocket_endpoint(ws: WebSocket):
    await ws.accept()
    ws_clients.add(ws)
    try:
        await ws.send_text(json.dumps({
            "type": "init",
            "data": await get_status(),
        }, ensure_ascii=False))
        while True:
            await ws.receive_text()
    except WebSocketDisconnect:
        ws_clients.discard(ws)
    except Exception:
        ws_clients.discard(ws)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host=SERVER_HOST, port=SERVER_PORT)
